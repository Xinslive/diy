"""窄 SDK：只返回序列化结果，复用宿主账套/角色/凭证/LedgerEngine 护栏。"""
from __future__ import annotations

from contextlib import contextmanager
from contextvars import ContextVar
from datetime import date, datetime
from decimal import Decimal
from types import SimpleNamespace
from fastapi import HTTPException
from pydantic import ValidationError
from sqlalchemy import MetaData, Table, func, or_, select
from sqlalchemy.exc import SQLAlchemyError
from app.core.db import SessionLocal
from app.core.deps import WRITE_ROLES, get_my_book, require_book_write
from app.models.book import Book
from app.models.system import SysUser, SystemSetting
from app.plugins import audit
from app.plugins.access import require_plugin_access
# ctx 暴露面的**登记表**住在纯数据模块里（worker 进程要读它，而 worker
# 不许 import 这里的数据库）。这里只做一次再导出，见文件末尾的说明。
from app.plugins.ctx_surface import CTX_METHOD_SERIALIZATION

# 只由宿主 ASGI 门禁设置，插件不能选择调用者。
current_principal = ContextVar("plugin_principal", default=None)
current_request_id = ContextVar("plugin_request_id", default="")

#: 当前请求里**已经解析好的调用者**。
#:
#: **按请求一份，不是按能力调用一次。**
#:
#: 为什么：宿主网关在**每个请求入口**已经做过一次 ``require_plugin_access``
#: （``loader.py`` 的 ``guard()``）。请求内部每一次能力调用又从头解析一遍
#: —— 那是**重复**，不是「更严」。一次解析要 7 条 SQL
#: （sys_user + plugin_registry + plugin_user_access + plugin_cohort +
#: plugin_worker_runtime×3）；实测一次共享卡工作台请求里解析了 **92 次**，
#: 合计 644 条纯授权样板，占那条请求全部 SQL 的 **57%**。
#:
#: **默认值是 None（也就是不缓存）**：没有进入 :func:`principal_scope` 的
#: 调用路径保持原样。这样「忘了开区间」只会慢，**不会**拿到上一个请求的调用者。
_principal_cache: ContextVar[dict | None] = ContextVar("plugin_principal_cache", default=None)


@contextmanager
def principal_scope():
    """进入一段「调用者按请求解析一次」的区间。

    **谁该开这个区间**：宿主里每一个「代表一个用户去调插件能力」的入口 ——
    HTTP 网关（``loader.py``）、插件钩子（``hooks.py``）、
    插件定时任务（``scheduler/service.py``）。三处都要开：
    少开一处，那条路径就退回「每次调用解析一次」（慢，但**正确**）。

    **区间必须与「一个用户的一次操作」一一对应**：
    区间里缓存的调用者，只在这一个操作内有效。
    """
    token = _principal_cache.set({})
    try:
        yield
    finally:
        _principal_cache.reset(token)


def request_user(db, pid):
    value = current_principal.get()
    if value is None or value[0] != pid:
        raise HTTPException(403, "插件能力只能在已鉴权请求内调用")
    cache = _principal_cache.get()
    # 缓存键**必须含 pid + user_id + generation**三样。
    #
    # 只拿 request_id 当键**不行**：钩子路径（``hooks.py``）**不设**
    # request_id，那个键会退化成「跨请求共享」—— 换个用户就会拿到上一个人的
    # 调用者。带上 pid 与 user_id 之后，键就绑死在「哪个插件的哪个用户」上，
    # 换人也换不到别人的。
    key = (pid, value[1], value[2])
    if cache is not None:
        hit = cache.get(key)
        if hit is not None:
            return hit
    user = db.get(SysUser, value[1])
    if not user or not user.active:
        raise HTTPException(401, "用户已停用")
    require_plugin_access(db, pid, user, generation=value[2])
    if cache is not None:
        # **把这个实例从会话里摘出来**：它会被本请求后面每一次能力调用共用，
        # 而那些调用各自开自己的会话。留在原会话里，它会跟着那一次 close
        # 变成「游离对象」；摘出来之后它就是一个**属性已加载**的独立对象。
        # 用到它的只有 id / username / active / is_admin 四个已加载的列，
        # 摘出来不影响任何一个。
        db.expunge(user)
        cache[key] = user
    return user


def _iso_day(value, what: str) -> date:
    """把一个日期参数解析成 ``date``（只认 ISO 的 ``YYYY-MM-DD``）。

    **为什么自己解析而不是用 ``date.fromisoformat`` 就完事**：调用方给的是
    字符串（插件那边是 JSON），给错了要**说清是哪个参数错了**，而不是抛一个
    ``ValueError`` 让插件收到一句看不懂的栈。
    """
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if not isinstance(value, str) or not value.strip():
        raise HTTPException(400, "%s无效" % what)
    try:
        return date.fromisoformat(value.strip()[:10])
    except ValueError:
        raise HTTPException(400, "%s必须是 YYYY-MM-DD" % what)


def _window_days(value, default: int, what: str) -> int:
    """窗口天数：不给就用默认，给了必须是正整数。"""
    if value is None:
        return default
    if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
        raise HTTPException(400, "%s天数必须是正整数" % what)
    return int(value)


#: 受限过滤 DSL 支持的操作符。刻意没有「原始表达式」这一类入口。
FILTER_OPERATORS = ("eq", "ne", "lt", "lte", "gt", "gte", "in", "like", "is_null")
#: 过滤值允许的标量类型：插件只能按值过滤，不能拼 SQL。
FILTER_SCALARS = (str, int, float, bool, Decimal, datetime, date)


#: 插件表的**反射结果缓存**：``(插件 id, 表名) -> Table``。
#:
#: **为什么缓存**：反射（``autoload``）是**每次读表都要重付一遍**的固定开销 ——
#: 实测（生产同款机器）每张表约 **21 ms**，而复用反射结果后的一次查询只要约
#: **3.5 ms**。插件一次请求要读几十次表（每本账一遍），不缓存就是几十次反射。
#:
#: **为什么这不是「省资源的小优化」**：插件默认在**宿主进程内**跑
#: （``PLUGIN_ISOLATION=off``），这些反射**同步地发生在 worker 的事件循环里** ——
#: 几十次反射就是几秒，卡住的是**整个宿主**（别的接口一起变慢），而不只是这个
#: 插件。共享卡对账插件一次列卡就要读三四十次表，正是这条路的重灾区。
#:
#: **什么时候会过期**：插件表的结构只由迁移改，而迁移**必须重启才生效**
#: （升级后 ``runtime_status`` 是 ``PENDING_RESTART``）—— 所以进程内的缓存与
#: 正在跑的代码永远是一致的。真要在运行期改结构（迁移提交那一刻），
#: 由 :func:`clear_plugin_table_cache` 显式失效。
_PLUGIN_TABLE_CACHE: dict[tuple[str, str], Table] = {}


def clear_plugin_table_cache(plugin_id: str | None = None) -> None:
    """清掉反射缓存。

    **迁移改了表结构之后调用**（:mod:`app.plugins.installer` 在迁移提交那一刻
    调它）—— 那一刻旧代码可能还在跑，旧代码看到的应该是**新**结构。
    """
    if plugin_id is None:
        _PLUGIN_TABLE_CACHE.clear()
        return
    for key in [item for item in _PLUGIN_TABLE_CACHE if item[0] == plugin_id]:
        _PLUGIN_TABLE_CACHE.pop(key, None)


def _reflected_table(plugin_id: str, name: str, db) -> Table:
    """插件表对象：**反射一次，之后复用**。

    反射出来的 ``Table`` 是**结构**（不是连接），拿它跨会话、跨线程查询都是
    安全的 —— SQLAlchemy 的 ``select(table)`` 只在拼语句时读它的列定义。
    """
    key = (plugin_id, name)
    table = _PLUGIN_TABLE_CACHE.get(key)
    if table is None:
        table = Table(name, MetaData(), autoload_with=db.connection(), schema="public")
        _PLUGIN_TABLE_CACHE[key] = table
    return table


def _statement_row_payload(row) -> dict:
    """一行流水 → 插件能存下来的形状。

    **金额不转字符串**：这里保住 ``Decimal``，让 IPC 那层的带类型标记把类型
    原样送回插件（见 ``app.plugins.ipc.values``）。插件自己要往外发 JSON 时
    再转 —— 转换只发生一次，且发生在他自己那一层。
    """
    return {
        "seq": row.seq,
        "tx_date": row.tx_date.isoformat() if row.tx_date else "",
        "tx_time": row.tx_time or "",
        "direction": row.direction,
        "amount": row.amount,
        "balance": row.balance,
        "summary": row.summary or "",
        "purpose": row.purpose or "",
        "counterparty_name": row.counterparty_name or "",
        "counterparty_account": row.counterparty_account or "",
        "extra": dict(row.extra or {}),
        "problems": list(row.problems or []),
    }


def _statement_meta_payload(meta) -> dict:
    """文件自报的账户与期间信息。"""
    return {
        "account_no": meta.account_no or "",
        "account_name": meta.account_name or "",
        "bank_name": meta.bank_name or "",
        "currency": meta.currency or "",
        "period_start": meta.period_start.isoformat() if meta.period_start else "",
        "period_end": meta.period_end.isoformat() if meta.period_end else "",
        "closing_balance": meta.closing_balance,
        "declared_income_count": meta.declared_income_count,
        "declared_income_total": meta.declared_income_total,
        "declared_expense_count": meta.declared_expense_count,
        "declared_expense_total": meta.declared_expense_total,
        "declared_total_count": meta.declared_total_count,
    }


def _filter_scalar(value):
    if value is None or isinstance(value, FILTER_SCALARS):
        return value
    raise HTTPException(400, "插件过滤值只能是标量")


def _filter_clauses(table, where):
    """把 ``where`` 翻译成 SQLAlchemy 条件。

    字段必须是本表真实列，值只能是标量；账套范围由调用方另行附加，
    过滤条件**改不掉**它。
    """
    if where is None:
        return []
    if not isinstance(where, dict):
        raise HTTPException(400, "插件过滤条件必须是对象")
    clauses = []
    for field, condition in where.items():
        if not isinstance(field, str) or field not in table.c:
            raise HTTPException(400, "插件过滤字段不存在")
        column = table.c[field]
        if isinstance(condition, dict):
            if not condition:
                raise HTTPException(400, "插件过滤条件不能为空")
            for operator, operand in condition.items():
                if operator not in FILTER_OPERATORS:
                    raise HTTPException(400, "插件过滤操作符不受支持")
                if operator == "in":
                    if not isinstance(operand, (list, tuple)) or not operand or len(operand) > 200:
                        raise HTTPException(400, "插件过滤集合无效")
                    clauses.append(column.in_([_filter_scalar(item) for item in operand]))
                elif operator == "like":
                    if not isinstance(operand, str) or len(operand) > 200:
                        raise HTTPException(400, "插件模糊匹配条件无效")
                    clauses.append(column.like(operand))
                elif operator == "is_null":
                    if not isinstance(operand, bool):
                        raise HTTPException(400, "插件空值判断条件无效")
                    clauses.append(column.is_(None) if operand else column.isnot(None))
                elif operator == "ne":
                    clauses.append(column != _filter_scalar(operand))
                elif operator == "lt":
                    clauses.append(column < _filter_scalar(operand))
                elif operator == "lte":
                    clauses.append(column <= _filter_scalar(operand))
                elif operator == "gt":
                    clauses.append(column > _filter_scalar(operand))
                elif operator == "gte":
                    clauses.append(column >= _filter_scalar(operand))
                else:
                    clauses.append(column == _filter_scalar(operand))
        else:
            clauses.append(column == _filter_scalar(condition))
    return clauses


def _order_clauses(table, order_by):
    if order_by is None:
        return list(table.primary_key.columns)
    if isinstance(order_by, str):
        order_by = [order_by]
    if not isinstance(order_by, (list, tuple)) or not order_by:
        raise HTTPException(400, "插件排序字段无效")
    columns = []
    for item in order_by:
        if not isinstance(item, str):
            raise HTTPException(400, "插件排序字段无效")
        field, _, direction = item.strip().partition(" ")
        if field not in table.c:
            raise HTTPException(400, "插件排序字段不存在")
        direction = (direction or "asc").strip().lower()
        if direction not in {"asc", "desc"}:
            raise HTTPException(400, "插件排序方向无效")
        columns.append(table.c[field].desc() if direction == "desc" else table.c[field].asc())
    return columns



def _book_scope_clause(table, user, book_id, across_books):
    """账套表的范围条件 —— **范围永远由宿主定**。

    单账套：``book_id == ?``；跨账套：``book_id IN (该用户有授权的账套)``。

    **为什么跨账套用子查询而不是把 id 列表传进来**：插件要是能指定
    「要读哪几本账」，那它就能指定「要读哪几本账」—— 范围就成了插件的选择，
    授权检查也就没了。子查询把集合锁死在 ``UserBook`` 上：管理员**也没有**
    隐含的全账套权限，能读到的只有他被授权的那些本。

    插件的 ``where`` 只能在这个范围**之上**再收窄（AND），收不宽。
    """
    if across_books:
        from app.models.system import UserBook
        return table.c.book_id.in_(
            select(UserBook.book_id).where(UserBook.user_id == user.id))
    return table.c.book_id == book_id


def context_for(manifest, *, plugin=False):
    """给一份 manifest 造它的能力上下文。

    ``plugin=True`` 是**插件拿到的那一份**（``loader.load_one`` 里交给
    ``register(registrar, ctx)`` 的就是它）；默认的 ``plugin=False`` 是**宿主
    自己**用的那一份（``plugins/services.py`` 拿它挂宿主路由）。

    两者差一处，但差的是要紧的一处：**插件面不含导入批次的「确认」能力。**
    确认只能走一次独立的 HTTP 请求（宿主路由
    ``POST /api/plugins/{id}/_host/imports/{ident}/confirm``），于是「用户点了
    确认」成为不可省略、且在审计里留痕的动作。

    **审批流不是安全边界，权限才是。** 宿主级校验（账套授权、批次归属、
    状态机、预览摘要一致、SAVEPOINT 原子提交）才是权威的那一层：即便插件
    绕过确认，数据完整性也守得住。插件侧确认破坏的是**审批语义** —— 一次
    请求内自己预览、自己确认、自己提交，用户那次独立的点击从来没有发生过，
    而审计里却会留下「用户确认过」的假象。
    """
    pid = manifest["id"]
    permissions = set(manifest["permissions"])
    declarations = {t["name"]: t["scope"] for t in manifest.get("tables", [])}

    def principal(db):
        return request_user(db, pid)

    def book_read(book_id: int):
        with SessionLocal() as db:
            user = principal(db)
            book = get_my_book(book_id, db, user)
            return {"id": book.id, "name": book.name}

    def book_list():
        """返回当前用户可访问的账套，附带能否写入的明确结果。

        插件是用户级授权，不绑定某一本账。文件导入插件需要在自己的页面
        让用户选择目标账套，因此不能只依赖宿主顶部当前账套。列表只来自
        当前用户的 ``UserBook`` 授权，``can_write`` 仅是提示；提交时仍由
        ``book_write`` 或凭证写入助手重新做实时校验。
        """
        from app.models.system import UserBook
        with SessionLocal() as db:
            user = principal(db)
            rows = db.execute(
                select(Book, UserBook.role)
                .join(UserBook, UserBook.book_id == Book.id)
                .where(UserBook.user_id == user.id)
                .order_by(UserBook.sort_order.asc().nulls_last(), Book.company_name, Book.name)
            ).all()
            return [
                {"id": book.id, "name": book.name, "company_name": book.company_name,
                 "status": book.status, "role": role,
                 "can_write": (role or "").strip().upper() in WRITE_ROLES}
                for book, role in rows
            ]

    def book_write(book_id: int):
        with SessionLocal() as db:
            user = principal(db)
            book = get_my_book(book_id, db, user)
            require_book_write(book_id, db, user)
            return {"id": book.id, "name": book.name}

    def upload_read(file, *, limit=None):
        """受宿主容量闸保护地读取插件路由收到的 UploadFile。

        插件可以决定更小的单次上限，但不能放宽宿主 ``import_max_bytes``。
        原始字节只在当前请求内返回给插件；持久化、去重和批次关联属于后续
        导入批次能力，不由这个低层读取助手偷偷完成。
        """
        from app.plugins.uploads import read_file

        with SessionLocal() as db:
            principal(db)
        return read_file(file, limit=limit)

    def ledger_read(book_id: int, period_id: int, *, include_unposted: bool = False):
        """某账套某期的科目余额表。

        ``include_unposted`` 是给「对账」这件事用的：宿主
        ``bankstatement.service.reconcile_balance`` 取账面余额时用的是
        ``LedgerEngine.for_ledger(..., include_unposted=True)`` ——
        **已制证但未过账的凭证也算在账面余额里**，否则差额里会混着
        「已经记了、只是没过账」这种噪声，让对账结论偏掉。

        默认 ``False``（与旧行为一致），老插件不受影响。
        """
        from app.api.ledger import balance
        with SessionLocal() as db:
            user = principal(db)
            book = get_my_book(book_id, db, user)
            return balance(book_id=book_id, period_id=period_id, from_period_id=None,
                           book=book, db=db, include_unposted=bool(include_unposted))

    def detail_ledger_read(book_id: int, subject_codes: list[str],
                           from_period_id: int, to_period_id: int, *,
                           include_unposted: bool = True):
        """某账套某段期间的**明细账**（逐笔 + 期初 + 本月合计 + 累计）。

        **为什么宿主出这条能力**：明细账的「余额怎么滚」不是加加减减 ——
        期初余额来自**引擎**（含结转、含跨年结转），逐月合计与本年累计的
        口径也只有一套规矩（见 ``app.ledger.queries._detail_account_rows``）。
        插件读不到宿主的业务表，也**不该自己再写一份**滚余额的算法 ——
        屏幕上的明细账与插件导出的明细账必须**是同一本账**。

        * ``subject_codes`` —— 科目**编码**清单。**按前缀匹配**（与
          :func:`entry_match` 同一口径）：登记 ``1002`` 也能取到记在
          ``100203`` 上的分录（那本账的科目表里没有 ``1002`` 这个末级时，
          取它的下级 —— 屏幕上查父科目看到的正是那一段）。
        * ``from_period_id`` / ``to_period_id`` —— 期间**区间**（含两端，
          按期间序号比大小，不按年月字符串比）。

        返回 ``{"subjects": [{"code", "ok", "error", "subject", "rows", ...}]}``。
        **每个科目单独成败**：一个科目不存在（这本账的科目表里没有它）不该
        让其余科目一起拿不到数 —— 返回里那一项带 ``ok=False`` 与原因，
        别的项照常给出行。
        """
        from app.ledger import queries
        from app.ledger.engine import LedgerEngine
        from app.models.accounting import Subject
        from app.models.book import Period

        if not isinstance(subject_codes, list) or not subject_codes:
            raise HTTPException(400, "科目编码清单不能为空")
        codes: list[str] = []
        for item in subject_codes:
            if not isinstance(item, str) or not item.strip():
                raise HTTPException(400, "科目编码无效")
            code = item.strip()
            if code not in codes:
                codes.append(code)

        with SessionLocal() as db:
            user = principal(db)
            book = get_my_book(book_id, db, user)
            to_period = db.get(Period, to_period_id)
            from_period = db.get(Period, from_period_id)
            if to_period is None or to_period.book_id != book.id:
                raise HTTPException(404, "止期不存在或不属于这本账")
            if from_period is None or from_period.book_id != book.id:
                raise HTTPException(404, "起期不存在或不属于这本账")
            if from_period.ord > to_period.ord:
                raise HTTPException(400, "起始期间不能晚于结束期间")
            # 科目表每个年度各有一份 —— 按**止期那一年的**科目表解析编码。
            year = to_period.year
            subjects = db.scalars(
                select(Subject)
                .where(Subject.book_id == book.id, Subject.year == year)
                .order_by(Subject.code)
            ).all()
            by_code = {item.code: item for item in subjects}
            # **一个引擎算到底**：逐个科目重建引擎，期初聚合就白算 N 遍
            # （与宿主导出整册明细账同一个理由）。
            eng = LedgerEngine.for_ledger(db, book.id, bool(include_unposted))
            out = []
            for code in codes:
                target = by_code.get(code)
                if target is None:
                    # 前缀匹配：登记的是父科目（1002），实际记在明细上（100203）。
                    children = [item for item in subjects if item.code.startswith(code)]
                    target = min(children, key=lambda item: (len(item.code), item.code)) \
                        if children else None
                if target is None:
                    out.append({"code": code, "ok": False,
                                "error": "这本账 %s 年的科目表里没有科目 %s" % (year, code)})
                    continue
                try:
                    payload = queries.detail_ledger(
                        db, book.id, target.id, from_period, to_period,
                        bool(include_unposted), eng=eng)
                except HTTPException as exc:
                    out.append({"code": code, "ok": False, "error": str(exc.detail)})
                    continue
                out.append({
                    "code": code, "ok": True, "error": "",
                    "subject": payload.get("subject") or {},
                    "from_period": payload.get("from_period") or "",
                    "to_period": payload.get("to_period") or "",
                    "rows": payload.get("rows") or [],
                })
        return {"subjects": out}

    def period_list(book_id: int):
        """账套的会计期间列表（id / 年月 / 账套内序号 / 是否已结账）。

        **为什么插件需要它**：插件算「这张卡在各账套的账面余额合计」时要知道
        取哪一期；跨账套的账期**可以不一样**（有的账套用 4-4-5 财年），
        所以不能拿宿主那份对账单的期间号直接套到别的账套上。

        插件**读不到宿主的 period 表**（没有这项能力），而凭证落哪一期是宿主
        按凭证日期在**那本账**里定的 —— 两边要能对上，就得有这份清单。
        """
        from app.models.book import Period
        with SessionLocal() as db:
            user = principal(db)
            book = get_my_book(book_id, db, user)
            rows = db.scalars(select(Period).where(
                Period.book_id == book.id).order_by(Period.ord)).all()
            return [{"id": row.id, "year": row.year, "month": row.month, "ord": row.ord,
                     "locked": bool(row.locked),
                     "closed_at": row.closed_at.isoformat() if row.closed_at else ""}
                    for row in rows]

    def statement_parse(data: bytes, *, book_id=None, column_mapping=None):
        """把一份对账单文件解析成结构化流水行 —— **不写宿主任何业务表**。

        这是给「**一张卡被多个账套共用**」那类场景的。那种对账单**不属于
        任何账套**：宿主 ``bank_statement`` 表按 (账套, 科目, 期间) 唯一，
        把一张卡绑进某个账套，会让每个账套都拥有一份完整流水（同一笔钱被
        记 N 次）、宿主余额对账报红、其余账套满屏假「存疑凭证」。

        所以共享卡的对账单是**插件实体**，宿主能贡献的只有**解析器** ——
        列识别、方向判定、自检、PDF 的两路解析与 OCR 兜底，这些在宿主
        里有 15 份真实银行样例对着金标准验过。插件自己重写一份，就是第二套
        口径：同一件事有两个答案。

        * ``column_mapping`` —— 列名 → 列号的映射，由插件自己存（共享卡跨
          账套，宿主那份画像按账套隔离，用不上）。传了就**整体替换**自动
          映射，不做逐字段合并（与宿主画像的语义一致）。
        * ``book_id`` —— 给了才落原件，且要检查当前用户能访问那本账。
          原件落在 ``data/statements/shared/``：它不属于任何账套，放进哪个
          账套的目录都是撒谎。

        返回的诊断、``via_ocr`` / ``via_merge`` 标志要**原样带给用户** ——
        OCR 是识别（会认错字），而且正文离开过本机，这是用户有权知道的事。
        """
        import hashlib

        from app.bankstatement import formats as bs_formats
        from app.bankstatement import service as bs_service
        from app.bankstatement import storage as bs_storage

        if not isinstance(data, (bytes, bytearray)) or not data:
            raise HTTPException(400, "对账单文件内容无效")
        mapping = None
        if column_mapping is not None:
            if not isinstance(column_mapping, dict):
                raise HTTPException(400, "列映射必须是对象")
            mapping = {}
            for key, value in column_mapping.items():
                if not isinstance(key, str) or not key or len(key) > 40:
                    raise HTTPException(400, "列映射的字段名无效")
                if type(value) is not int or value < 0:
                    raise HTTPException(400, "列映射的列号无效")
                mapping[key] = value
        raw = bytes(data)
        with SessionLocal() as db:
            user = principal(db)
            owner_book = None
            if book_id is not None:
                owner_book = get_my_book(book_id, db, user)
            try:
                parsed = bs_service.parse_only(db, raw, column_mapping=mapping)
            except ValueError as exc:
                raise HTTPException(400, f"这份对账单解析不出来：{exc}") from exc
            period_start, period_end = bs_service.effective_period(parsed)
            file_info = {"relative_path": "", "container": bs_formats.sniff(raw),
                         "size": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}
            if owner_book is not None:
                file_info["relative_path"] = bs_storage.store_shared(file_info["container"], raw)
            return {
                "container": parsed.container,
                "sheet": parsed.sheet,
                "header_row": parsed.header_row,
                "header_cells": list(parsed.header_cells),
                "column_mapping": dict(parsed.column_mapping),
                "fingerprint": parsed.fingerprint,
                "rows": [_statement_row_payload(row) for row in parsed.rows],
                "meta": _statement_meta_payload(parsed.meta),
                "diagnosis": list(parsed.diagnosis),
                "via_ocr": bool(parsed.via_ocr),
                "via_merge": bool(parsed.via_merge),
                "merge_conflicts": list(parsed.merge_conflicts),
                "period_start": period_start.isoformat() if period_start else "",
                "period_end": period_end.isoformat() if period_end else "",
                "income_total": parsed.income_total,
                "expense_total": parsed.expense_total,
                "file": file_info,
                "owner_book_id": owner_book.id if owner_book is not None else None,
            }

    def statement_read(book_id: int):
        """该账套在宿主已绑的对账单**清单**。

        **只给清单，不给流水行内容**：插件要的是「这张卡是不是已经在这个
        账套里绑过了」这一个事实 —— 如果绑过，插件再制一遍凭证就是**双重
        记账**（钱在账上多记一次）。行的内容各算各的，混在一起反而说不清
        哪一行归谁。

        科目码取**活科目**，没有活科目才回落快照列 —— 与宿主流水对账页面
        同一份判据（``caliber.account_subject_code``）：只差一位的编码
        （1001 与 100201）肉眼扫过去看不出来，两边不一致比哪边旧更糟。
        """
        from app.bankstatement.caliber import account_subject_code
        from app.models.accounting import Subject
        from app.models.bankstatement import BankStatement
        with SessionLocal() as db:
            user = principal(db)
            book = get_my_book(book_id, db, user)
            statements = db.scalars(select(BankStatement).where(
                BankStatement.book_id == book.id).order_by(
                BankStatement.period_start.desc().nullslast(),
                BankStatement.id.desc())).all()
            result = []
            for statement in statements:
                subject = db.get(Subject, statement.account_subject_id)
                result.append({
                    "id": statement.id,
                    "account_subject_id": statement.account_subject_id,
                    "account_subject_code": account_subject_code(subject, statement),
                    "account_subject_name": subject.name if subject else "",
                    "filename": statement.filename,
                    "bank_name": statement.bank_name,
                    "account_no": statement.account_no,
                    "period_start": (statement.period_start.isoformat()
                                     if statement.period_start else ""),
                    "period_end": (statement.period_end.isoformat()
                                   if statement.period_end else ""),
                    "closing_balance": (str(statement.closing_balance)
                                         if statement.closing_balance is not None else ""),
                    "row_count": statement.row_count,
                    "parse_status": statement.parse_status,
                })
            return result

    def subject_read(book_id: int, year: int | None = None):
        """当前账套**某个年度**的科目表。插件用它把文件里的科目**编码**
        换成凭证要的 ``subject_id``。

        返回**全部**科目（末级与非末级、启用与停用都给），而不是像
        ``/subjects/flat`` 那样默认只给末级且启用的：插件要区分「这个编码根本
        不存在」「这个科目已停用」「这是个上级科目，不能记账」三种情况才能
        给人报出有用的错。只给可用集合的话，插件只能说「找不到」，用户拿到
        一句没有指向的错。筛选是插件自己的事，宿主不替它做这个判断。

        **科目表每个年度各有一份**（迁移 0032）：2026 的 1002 和 2027 的
        1002 是两行不同的记录、两个 id。所以要指定年度 —— 不指定就用
        **账套当前在记的那一年**。插件要把凭证记在哪一年，就得取哪一年的表，
        否则拿到的 id 属于另一年，凭证那一步会被库层挡下来。
        """
        from app.accounting.years import resolve_year
        from app.models.accounting import Subject
        with SessionLocal() as db:
            user = principal(db)
            book = get_my_book(book_id, db, user)
            resolved = resolve_year(db, book, year)
            rows = db.scalars(
                select(Subject)
                .where(Subject.book_id == book.id, Subject.year == resolved)
                .order_by(Subject.code)
            ).all()
            return [{"id": s.id, "code": s.code, "name": s.name, "level": s.level,
                     "category": s.category, "direction": s.direction,
                     "is_leaf": s.is_leaf, "active": s.active, "year": s.year}
                    for s in rows]

    def entry_match(book_id: int, subject_codes: list[str], rows: list[dict], *,
                    date_from: str | None = None, date_to: str | None = None,
                    primary_days: int | None = None, review_days: int | None = None):
        """把一批流水行与账上**现金科目**的分录一一配对。

        **为什么宿主出这条能力，而不是让插件自己查分录**：插件读不到宿主的
        业务表（``data.read`` 只能读插件自己声明的表 —— ``table_access`` 对
        没声明的表直接 403）。而「一行流水对应账上哪张凭证」这件事**宿主的
        流水对账那条线早就在做**：判据（金额相等 + 现金侧方向对应 + 日期
        窗口）、一对一分配、跨期「待复核」那一档，全在
        :mod:`app.bankstatement.allocate` 里，并且是那一条线**唯一的一把尺子**
        （那个模块的文档写得很直白：两页各用各的数，正是「两页都不报」的
        成因）。插件自己再写一份，就是第二套口径 —— 同一件事有两个答案。

        所以这里**不重新实现配对**，只把宿主那把尺子接到插件面上：插件给出行
        （日期 + 金额 + 方向），宿主按自己的口径配好对回来。

        * ``subject_codes`` —— 参与配对的**科目编码**（通常是这张卡在各账套
          登记的现金科目，如 ``100203``）。按**前缀**匹配：登记 ``1002`` 也能
          配到记在 ``100203`` 上的分录（父科目登记、明细记账是常事）。
        * ``rows`` —— ``[{"id": 1, "tx_date": "2026-09-02",
          "amount": "2829.58", "direction": "IN"}]``。``direction`` 用**流水
          方向**（``IN`` 收钱 / ``OUT`` 付钱）；现金侧方向由宿主换算 ——
          银行账与会计凭证的方向是相反的（收钱 = 借方），换算写在宿主这一侧，
          插件不用自己记这张表。

        返回 ``{"pairs": [...], "unmatched_row_ids": [...]}``。每条 pair 带
        ``days_diff`` 与 ``in_primary``（日期差是否在主窗口内）—— 后者就是宿主
        那条线的「已入账 / 待复核」两档：**在主窗口内是同一笔，出了窗口只是
        金额对得上，要人判**。

        **窗口可以由调用方给**（1.4.0 起）—— 这一段是给「对账单流水是本月、
        记账在上月」那种错位场景的：

        * ``date_from`` / ``date_to`` —— **分录的扫描区间**（含两端）。不给就是
          老口径：按流水行的日期前后各留 30 天。那 30 天是宿主的默认值，
          **跨月错位够不着**：9 月 30 日的流水配 8 月 5 日的凭证差 56 天，
          压根不会进候选池 —— 看上去就是「账上没记这笔」。
        * ``primary_days`` —— **主窗口**（几天之内算「已入账」，默认 7）。
        * ``review_days`` —— **复核窗口**（几天之内还算「待复核」，默认 30）。

        **尺子还是宿主这把**（``allocate`` 一字未改），变的只是**窗口** ——
        窗口本来就是「同一笔」这一定义的一部分，而**窗口该多大有时不是
        算得出来的**：9 月的对账单去对 8 月的账，跨月这件事**用户知道**，
        算不出来。所以窗口交给调用方给，判定仍然只有一份口径。

        **给了 ``date_from``/``date_to`` 就不再看流水行的日期**：那正是
        「对账单的月份不是记账的月份」这件事的要点 —— 流水行的日期是**银行
        记账日**，分录的日期是**实际发生日**，两者对不上时，扫哪一段要由
        调用方说了算。
        """
        from datetime import timedelta

        from app.bankstatement.allocate import (
            PRIMARY_WINDOW_DAYS,
            REVIEW_WINDOW_DAYS,
            EntrySide,
            RowSide,
            allocate,
        )
        from app.models.accounting import Entry, Subject, Voucher

        if not isinstance(subject_codes, list) or not subject_codes:
            raise HTTPException(400, "科目编码清单不能为空")
        codes: list[str] = []
        for item in subject_codes:
            if not isinstance(item, str) or not item.strip():
                raise HTTPException(400, "科目编码无效")
            code = item.strip()
            if code not in codes:
                codes.append(code)
        if not isinstance(rows, list):
            raise HTTPException(400, "流水行清单格式无效")
        if len(rows) > 5000:
            raise HTTPException(400, "一次最多配对 5000 行")

        sides: list[RowSide] = []
        dated: list[date] = []
        for item in rows:
            if not isinstance(item, dict):
                raise HTTPException(400, "流水行格式无效")
            try:
                row_id = int(item.get("id"))
            except (TypeError, ValueError):
                raise HTTPException(400, "流水行缺少 id")
            amount = None
            raw_amount = item.get("amount")
            if raw_amount is not None and raw_amount != "":
                try:
                    amount = Decimal(str(raw_amount))
                except (ArithmeticError, TypeError, ValueError):
                    amount = None
            tx_date = None
            raw_date = item.get("tx_date")
            if isinstance(raw_date, date):
                tx_date = raw_date
            elif isinstance(raw_date, str) and raw_date.strip():
                try:
                    tx_date = date.fromisoformat(raw_date.strip()[:10])
                except ValueError:
                    tx_date = None
            try:
                seq = int(item.get("seq") or 0)
            except (TypeError, ValueError):
                seq = 0
            sides.append(RowSide(
                id=row_id, tx_date=tx_date, amount=amount,
                direction=str(item.get("direction") or "").strip().upper(), seq=seq))
            if tx_date is not None:
                dated.append(tx_date)

        if not dated:
            # **没有日期就配不了**：配对的第一件套就是日期窗口。
            # 认不出日期的行不是「没配上」，是「还没资格配上」—— 但对插件来说
            # 结果一样：没有候选。
            return {"pairs": [], "unmatched_row_ids": [side.id for side in sides]}

        # ---- 窗口：默认按流水行的日期前后各留 30 天（老口径） -------------
        #
        # 调用方给了区间就用它给的 —— 「流水是本月的、记账是上月的」那种
        # 错位场景里，**该扫哪一段由人说了算**（插件那边是用户指定的匹配
        # 月份范围），算不出来。
        lo = min(dated) - timedelta(days=REVIEW_WINDOW_DAYS)
        hi = max(dated) + timedelta(days=REVIEW_WINDOW_DAYS)
        if (date_from is None) != (date_to is None):
            raise HTTPException(400, "扫描区间的起止日期要一起给")
        if date_from is not None:
            lo = _iso_day(date_from, "扫描区间起日")
            hi = _iso_day(date_to, "扫描区间止日")
            if lo > hi:
                raise HTTPException(400, "扫描区间的起日不能晚于止日")
        primary_window = _window_days(primary_days, PRIMARY_WINDOW_DAYS, "主窗口")
        review_window = _window_days(review_days, REVIEW_WINDOW_DAYS, "复核窗口")
        if review_window < primary_window:
            raise HTTPException(400, "复核窗口不能小于主窗口")

        with SessionLocal() as db:
            user = principal(db)
            book = get_my_book(book_id, db, user)
            rows_found = db.execute(
                select(Entry, Voucher, Subject)
                .join(Voucher, Voucher.id == Entry.voucher_id)
                .join(Subject, Subject.id == Entry.subject_id)
                .where(
                    Voucher.book_id == book.id,
                    Entry.book_id == book.id,
                    or_(*[Subject.code.startswith(code) for code in codes]),
                    Voucher.voucher_date >= lo,
                    Voucher.voucher_date <= hi,
                )
            ).all()
            entry_sides = [
                EntrySide(
                    id=int(entry.id), voucher_id=int(entry.voucher_id),
                    voucher_no=voucher.voucher_no or "",
                    voucher_date=voucher.voucher_date,
                    amount=Decimal(str(entry.amount)), direction=entry.direction)
                for entry, voucher, _subject in rows_found
            ]

        result = allocate(sides, entry_sides, primary_window=primary_window,
                          review_window=review_window)
        pairs = [{
            "row_id": pair.row_id,
            "entry_id": pair.entry_id,
            "voucher_id": pair.voucher_id,
            "voucher_no": pair.voucher_no,
            "voucher_date": pair.voucher_date.isoformat(),
            "amount": str(pair.amount),
            "days_diff": pair.days_diff,
            "in_primary": bool(pair.in_primary),
        } for pair in result.pairs.values()]
        matched = {pair["row_id"] for pair in pairs}
        return {"pairs": pairs,
                "unmatched_row_ids": [side.id for side in sides if side.id not in matched]}

    def export_xlsx(filename: str, columns: list[dict], rows: list[dict],
                    sheet_name: str = "Sheet1"):
        """把一张表导成 Excel（``.xlsx``），回传 **base64**。

        **为什么在宿主生成，而不是插件自己造**：``.xlsx`` 是个 ZIP 包，而插件
        沙箱只放行 ``openpyxl.load_workbook``（**读**），不放行 ``Workbook``
        （**写**）—— 插件能解析用户上传的表格，但不能造一个新的。而导出是宿主
        本来就有的能力（凭证、账簿、余额表的导出都在这一侧），这里只是把
        「任意一张表」这一种也开放出来。

        **为什么回传 base64 而不是字节**：插件路由的返回值要能过 JSON（插件
        后端将来可能跑在独立进程里，那时返回值走 IPC 序列化）。base64 字符串
        两头都能用：前端 ``atob`` 一下就是 Blob。

        * ``columns`` —— ``[{"key": "seq", "title": "行", "width": 6}]``。
          ``width`` 是 Excel 列宽（字符数），不给就用默认。
        * ``rows`` —— 行列表，每行一个 dict，键与 ``columns[].key`` 对应。
          取不到的键**留空，不猜**。
        """
        import base64
        import io

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        if not isinstance(columns, list) or not columns:
            raise HTTPException(400, "导出列清单不能为空")
        if len(columns) > 64:
            raise HTTPException(400, "导出列不能超过 64 列")
        if not isinstance(rows, list):
            raise HTTPException(400, "导出行清单格式无效")
        if len(rows) > 20000:
            raise HTTPException(400, "一次最多导出 20000 行")

        keys: list[str] = []
        titles: list[str] = []
        widths: list[float] = []
        for column in columns:
            if not isinstance(column, dict):
                raise HTTPException(400, "导出列格式无效")
            key = column.get("key")
            title = column.get("title")
            if not isinstance(key, str) or not key:
                raise HTTPException(400, "导出列缺少 key")
            if not isinstance(title, str) or not title:
                raise HTTPException(400, "导出列缺少标题")
            keys.append(key)
            titles.append(title)
            width = column.get("width")
            widths.append(float(width) if isinstance(width, (int, float)) else 14.0)

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name if isinstance(sheet_name, str) and sheet_name.strip() else "Sheet1"
        ws.append(titles)
        fill = PatternFill("solid", fgColor="DCE6F1")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
        for row in rows:
            if not isinstance(row, dict):
                raise HTTPException(400, "导出行格式无效")
            ws.append([row.get(key, "") for key in keys])
        buf = io.BytesIO()
        wb.save(buf)
        name = filename if isinstance(filename, str) and filename.strip() else "导出.xlsx"
        if not name.lower().endswith(".xlsx"):
            name = name + ".xlsx"
        return {"filename": name, "content_base64": base64.b64encode(buf.getvalue()).decode("ascii")}

    def voucher_create(book_id: int, payload: dict):
        from app.api.schemas import VoucherIn
        from app.api.vouchers import create
        with SessionLocal() as db:
            user = principal(db)
            book = get_my_book(book_id, db, user)
            link = require_book_write(book_id, db, user)
            try:
                data = VoucherIn.model_validate(payload)
            except ValidationError as exc:
                raise HTTPException(422, "插件凭证输入不符合宿主校验规则") from exc
            return create(payload=data, book=book, db=db, _link=link, user=user)

    def voucher_batch_create(book_id: int, payloads: list[dict]):
        """一次建多张凭证，**整批同一个事务**。

        单张 ``voucher_create`` 每次自己开一个 session 并提交，所以拿它循环
        500 行：一行失败时前 499 张**已经落库**了。「按文件批量生成凭证」正是
        这个形状 —— 失败留下半批凭证，重跑又因为没有水位而重复，用户要自己去
        账上把半批一张张删掉。这里对齐宿主导入那条路（``importer.commit``）：
        逐张 ``commit=False``，退出时才统一提交，任一张抛错整个事务回滚。

        错误信息带上**第几张**，否则插件只能回一句「失败了」，用户面对 500 行
        不知道该改哪一行。
        """
        from app.accounting import vouchers as svc
        from app.api.schemas import VoucherIn
        from app.api.vouchers import _book_periods, _book_subjects, _serialize
        from app.core.config import settings

        if not isinstance(payloads, list) or not payloads:
            raise HTTPException(400, "批量创建至少需要一张凭证")
        limit = settings.import_max_vouchers
        if len(payloads) > limit:
            raise HTTPException(400, f"一次最多创建 {limit} 张凭证，本次传了 {len(payloads)} 张。"
                                     "请拆成多批，每批成功后再提交下一批")
        with SessionLocal.begin() as db:
            user = principal(db)
            book = get_my_book(book_id, db, user)
            require_book_write(book_id, db, user)
            made = []
            for i, raw in enumerate(payloads, 1):
                try:
                    data = VoucherIn.model_validate(raw)
                except ValidationError as exc:
                    raise HTTPException(422, f"第 {i} 张凭证输入不符合宿主校验规则：{exc.errors()[0]['msg']}") from exc
                try:
                    made.append(svc.create_voucher(
                        db, book_id=book.id, user=user, voucher_date=data.voucher_date,
                        entries=[e.model_dump() for e in data.entries], remark=data.remark,
                        attachments=data.attachments, source=data.source,
                        voucher_no=data.voucher_no, insert_after_id=data.insert_after,
                        commit=False))
                except (ValueError, ArithmeticError) as e:
                    raise HTTPException(400, f"第 {i} 张凭证没建成，整批已回滚：{e}") from e
            # 科目一次取回给所有凭证复用（见 _book_subjects 的 docstring）。
            subs = _book_subjects(db, book.id)
            # 期间同样一次取回（一条查询）：年度随凭证一起出去，形状与 REST 那边一致 ——
            # 同一个「凭证」两处形状不同，是个迟早会咬人的坑。
            periods = _book_periods(db, book.id, {v.period_id for v in made})
            return [_serialize(db, v, subs, periods) for v in made]

    def table_access(db, user, name, book_id, write=False, across_books=False):
        """插件能碰这张表的哪个范围。

        ``across_books`` —— **跨账套读**，只对账套表的**读**有效：
        范围是「当前用户有授权的全部账套」（``UserBook``），
        由宿主算好，插件既不能指定集合、也不能改写它。

        **为什么不允许跨账套写**：写必须点名账套 —— 账套级写权限、
        归属判断、审计留痕全都挂在「哪一本账」上。跨着写，权限就没有落点。
        """
        if name not in declarations:
            raise HTTPException(403, "只能访问当前插件声明的数据表")
        scope = declarations[name]
        if scope == "book":
            if across_books:
                if write:
                    raise HTTPException(400, "跨账套读取不能用于写入，写入必须点名账套")
                if book_id is not None:
                    raise HTTPException(400, "跨账套读取不能同时指定 book_id")
            else:
                if type(book_id) is not int:
                    raise HTTPException(400, "账套表需要 book_id")
                get_my_book(book_id, db, user)
                if write:
                    require_book_write(book_id, db, user)
        elif across_books:
            raise HTTPException(400, "全局表没有账套范围，不能跨账套读取")
        elif book_id is not None:
            raise HTTPException(400, "全局表不能指定 book_id")
        elif write and not user.is_admin:
            raise HTTPException(403, "全局配置写入需要管理员")
        return _reflected_table(pid, name, db)

    def data_read(name: str, *, book_id=None, limit=100, offset=0, where=None, order_by=None,
                  across_books=False):
        """读插件自有表的一页。

        旧契约只有「按主键顺序 + 最多 500 行」，长期运行的插件只能把整表读进
        内存自己过滤。这里补上**受限过滤**、**分页**和**计数**：

        - ``where``：``{"字段": 值}`` 或 ``{"字段": {"gte": 值}}``。字段必须是
          本表真实列，值只能是标量 —— 没有任何拼 SQL 的入口。
        - ``offset``/``limit``：分页；``limit`` 仍然夹在 1..500。
        - ``order_by``：``"字段"`` 或 ``"字段 desc"``，默认按主键。
        - ``across_books``：**跨账套读**（只对账套表有效，见 :func:`table_access`）。
          范围是「当前用户有授权的全部账套」，由宿主用 ``UserBook`` 子查询定死，
          插件改不掉。行里带各自的 ``book_id``。

        账套范围（book scope）永远由宿主附加，过滤条件改不掉它。
        返回值仍是行列表 —— 与旧契约一致，新增能力都是可选参数。

        **为什么需要跨账套读**：插件自有表里，一行落在**某一本账**里（账套级备份
        才覆盖得到）。跨账套的业务对象（一张共享卡横跨好几本账）于是只能
        「按账套读一遍再合并」—— 读几次就随账套数线性涨。宿主实测过：
        一次工作台请求读 12 本账就是 12 次往返（外加每份对账单再读一遍归属，
        变成 N×M 次），而那些表**总共只有一百多行**。
        """
        with SessionLocal() as db:
            user = principal(db)
            table = table_access(db, user, name, book_id, across_books=across_books)
            stmt = select(table)
            if declarations[name] == "book":
                stmt = stmt.where(_book_scope_clause(table, user, book_id, across_books))
            stmt = stmt.where(*_filter_clauses(table, where))
            stmt = stmt.order_by(*_order_clauses(table, order_by))
            stmt = stmt.offset(max(0, int(offset))).limit(max(1, min(int(limit), 500)))
            return [dict(row) for row in db.execute(stmt).mappings()]

    def data_count(name: str, *, book_id=None, where=None, across_books=False):
        """符合条件的行数。

        与 :func:`data_read` 用同一套受限过滤条件；单独提供而不是把
        ``data.read`` 的返回值改成字典，是为了不破坏既有插件的调用契约。
        """
        with SessionLocal() as db:
            user = principal(db)
            table = table_access(db, user, name, book_id, across_books=across_books)
            stmt = select(func.count()).select_from(table)
            if declarations[name] == "book":
                stmt = stmt.where(_book_scope_clause(table, user, book_id, across_books))
            stmt = stmt.where(*_filter_clauses(table, where))
            return int(db.scalar(stmt) or 0)

    def data_insert(name: str, values: dict, *, book_id=None):
        with SessionLocal.begin() as db:
            user = principal(db)
            table = table_access(db, user, name, book_id, True)
            if not isinstance(values, dict) or set(values) - set(table.c.keys()) or ("book_id" in values and values["book_id"] != book_id):
                raise HTTPException(400, "插件数据字段或账套不合法")
            record = dict(values)
            if declarations[name] == "book":
                record["book_id"] = book_id
            return dict(db.execute(table.insert().values(**record).returning(*table.c)).mappings().one())

    def data_insert_many(name: str, values_list: list, *, book_id=None):
        """批量插入插件自有表的行 —— **整批同一个事务**。

        与 :func:`voucher_batch_create` 是同一条理由：``data.insert`` 一次
        只插一行、一次一个事务，拿它循环 500 行，第 N 行失败时**前 N−1 行
        已经落库**了。一份对账单动辄上万行流水，逐行插入既慢，又留下半批
        脏数据 —— 用户看到的是「插了一半」，而他以为这是一次操作。

        错误信息要指到**第几行**，否则插件只能回一句「失败了」，用户面对
        上万行不知道该改哪一行。
        """
        if not isinstance(values_list, list) or not values_list:
            raise HTTPException(400, "批量插入至少需要一行")
        if len(values_list) > 20000:
            raise HTTPException(400, "一次最多插入 20000 行")
        with SessionLocal.begin() as db:
            user = principal(db)
            table = table_access(db, user, name, book_id, True)
            columns = set(table.c.keys())
            records = []
            for index, values in enumerate(values_list, 1):
                if not isinstance(values, dict) or not values \
                        or set(values) - columns \
                        or ("book_id" in values and values["book_id"] != book_id):
                    raise HTTPException(400, f"第 {index} 行的插件数据字段或账套不合法")
                record = dict(values)
                if declarations[name] == "book":
                    record["book_id"] = book_id
                records.append(record)
            try:
                inserted = db.execute(
                    table.insert().values(records).returning(*table.c)).mappings().all()
            except SQLAlchemyError as exc:
                raise HTTPException(400, f"批量插入失败，整批已回滚：{exc}") from exc
            return [dict(row) for row in inserted]

    def data_update(name: str, ident: int, values: dict, *, book_id=None):
        with SessionLocal.begin() as db:
            user = principal(db)
            table = table_access(db, user, name, book_id, True)
            if "id" not in table.c or not isinstance(values, dict) or not values or set(values) - set(table.c.keys()) or set(values) & {"id", "book_id"}:
                raise HTTPException(400, "插件更新字段不合法，不能修改主键或账套")
            stmt = table.update().where(table.c.id == ident)
            if declarations[name] == "book":
                stmt = stmt.where(table.c.book_id == book_id)
            return {"updated": db.execute(stmt.values(**values)).rowcount}

    def data_delete(name: str, ident: int, *, book_id=None):
        with SessionLocal.begin() as db:
            user = principal(db)
            table = table_access(db, user, name, book_id, True)
            if "id" not in table.c:
                raise HTTPException(400, "此助手要求 id 主键")
            stmt = table.delete().where(table.c.id == ident)
            if declarations[name] == "book":
                stmt = stmt.where(table.c.book_id == book_id)
            return {"deleted": db.execute(stmt).rowcount}

    def audit_write(action: str):
        if not isinstance(action, str) or not action or len(action) > 32:
            raise HTTPException(400, "审计事件标识无效")
        with SessionLocal.begin() as db:
            user = principal(db)
            audit.event(db, "event", pid, user=user, code=action)

    def settings_key(key):
        from app.plugins.manifest import table_prefix
        full = f"plugin.{table_prefix(pid)[7:19]}.{key}"
        if not isinstance(key, str) or not key.isascii() or not key.replace("_", "").isalnum() or len(full) > 64:
            raise HTTPException(400, "插件设置 key 无效或过长")
        return full

    def setting_get(key: str):
        with SessionLocal() as db:
            principal(db)
            setting = db.get(SystemSetting, settings_key(key))
            return setting.value if setting else None

    def setting_put(key: str, value: str):
        with SessionLocal.begin() as db:
            user = principal(db)
            if not user.is_admin:
                raise HTTPException(403, "全局配置写入需要管理员")
            if not isinstance(value, str) or len(value) > 4096:
                raise HTTPException(400, "插件设置值过长")
            ident = settings_key(key)
            item = db.get(SystemSetting, ident)
            if item:
                item.value = value
            else:
                db.add(SystemSetting(key=ident, value=value))

    exposed = {"plugin_id": pid}
    for permission, key, fn in [("book:read", "book_read", book_read), ("book:read", "book_list", book_list),
        ("book:write", "book_write", book_write),
        ("upload:read", "upload_read", upload_read),
        ("ledger:read", "ledger_read", ledger_read), ("subject:read", "subject_read", subject_read),
        # 明细账与余额表同属「读账」这一条权限 —— 不另开一条，免得权限表越滚越长。
        ("ledger:read", "detail_ledger_read", detail_ledger_read),
        ("book:read", "period_list", period_list),
        # 配对复用宿主那把尺子（``app.bankstatement.allocate``），插件不自己写第二份。
        ("entry:read", "entry_match", entry_match),
        # 导出在宿主生成：``.xlsx`` 要 openpyxl 的写侧，插件沙箱只放行读侧。
        ("export:write", "export_xlsx", export_xlsx),
        ("voucher:write", "voucher_create", voucher_create),
        ("audit:write", "audit_write", audit_write)]:
        if permission in permissions:
            exposed[key] = fn
    # 批量与单张同权限：批量只是把同一套校验放进一个事务，不额外给权。
    if "voucher:write" in permissions:
        exposed["voucher_batch_create"] = voucher_batch_create
    # 对账单两项能力都**只读宿主**：解析不落宿主业务表，清单不含流水行内容。
    if "statement:parse" in permissions or "statement:read" in permissions:
        statement_ns = SimpleNamespace()
        if "statement:parse" in permissions:
            statement_ns.parse = statement_parse
        if "statement:read" in permissions:
            statement_ns.read = statement_read
        exposed["statement"] = statement_ns
    if "plugin_data:rw" in permissions:
        from app.plugins.configuration import service as configuration_service
        exposed["config"] = configuration_service(manifest, principal)
        exposed["data"] = SimpleNamespace(read=data_read, count=data_count, insert=data_insert,
                                          insert_many=data_insert_many,
                                          update=data_update, delete=data_delete)
    if "settings:rw" in permissions:
        exposed["settings"] = SimpleNamespace(get=setting_get, put=setting_put)
    if "import:rw" in permissions:
        from app.plugins.imports import service
        # 插件面（plugin=True）**不含 confirm**：确认必须是一次独立的请求。
        exposed["imports"] = service(manifest, principal, expose_confirm=not plugin)
    return SimpleNamespace(**exposed)


# ---------------------------------------------------------------------------
# ctx 表面：每个方法的返回值**能不能过 JSON**
# ---------------------------------------------------------------------------
# 权威实现在 :mod:`app.plugins.ctx_surface`（纯数据模块，worker 也要读）；
# 上面已 import 过来，这里只保留**为什么**它不住在这个模块里的说明。
