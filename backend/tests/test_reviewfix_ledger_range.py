"""跨 3 个以上年度的区间查询：中间年度的发生额不许被丢掉（审查修复）。

``queries.subjects_for_periods`` 的 docstring 写着「区间涉及哪几个年度，
就取哪几个年度的科目表（并集）」，而调用方（余额表接口、打印、导出）
传的是 ``[from_period, period]`` **两个**期间 —— 于是只算得出**首尾两年**。

中间年度（2027）的科目表压根没进查询，而 2027 年的分录挂在**2027 那套
科目行**上（另一个 id）：``merge_by_code`` 拿首尾两年的 id 去聚合，
2027 那半截**一条都匹配不到** —— 区间借方合计凭空少一整年，
而行内仍然自洽（期初 + 发生额 = 期末），**没有任何地方会报错**。

修法：取表按**区间内的全部期间**取；``merge_by_code`` 再加一道
fail-closed 断言（年度没取全就报错，而不是默默算少）。
"""
from __future__ import annotations

from decimal import Decimal

import pytest
from sqlalchemy import select

from app.models import Period


def _period(db, book, year: int, month: int) -> Period:
    p = db.scalar(select(Period).where(
        Period.book_id == book.id, Period.year == year, Period.month == month))
    assert p is not None, f"账套 {book.id} 没有 {year}-{month:02d} 这个期间"
    return p


def _open_years(db, book, years: list[int]) -> None:
    """铺出这些年度的期间与各自的科目表。"""
    from app.accounting.periods import ensure_periods
    from app.accounting.years import copy_chart

    for year in years:
        copy_chart(db, book.id, year - 1, year)
        ensure_periods(db, book, year)
    db.commit()


def _balance_row(rows: list[dict], code: str) -> dict:
    row = next((r for r in rows if r["code"] == code), None)
    assert row is not None, f"余额表里没有 {code} 这一行"
    return row


def test_跨三年区间的借方合计包含中间年度(client, auth, db, book, make_voucher):
    """2026-01 ~ 2028-12 的区间借方必须包含 2027 那笔。

    修好之前：区间只取首尾两年的科目表 → 2027 的 200 元**整段消失**，
    而返回的表仍然是自洽的、不会报错。
    """
    _open_years(db, book, [2027, 2028])

    make_voucher("2026-06-10", dr=[("1002", 100, "x")], cr=[("5001", 100, "x")],
                 posted=True)
    make_voucher("2027-06-10", dr=[("1002", 200, "x")], cr=[("5001", 200, "x")],
                 posted=True)
    make_voucher("2028-06-10", dr=[("1002", 400, "x")], cr=[("5001", 400, "x")],
                 posted=True)
    db.commit()

    first = _period(db, book, 2026, 1)
    last = _period(db, book, 2028, 12)
    r = client.get(f"/api/books/{book.id}/ledger/balance", headers=auth, params={
        "period_id": last.id, "from_period_id": first.id})
    assert r.status_code == 200, r.text[:300]

    row = _balance_row(r.json()["rows"], "1002")
    total = Decimal(str(row["period_debit"]))
    assert total == Decimal("700"), (
        f"区间借方合计应当是 100 + 200 + 400 = 700，实际 {total} —— "
        f"中间年度（2027）的发生额被整段丢掉了")


def test_跨三年区间导出的合计与页面一致(client, auth, db, book, make_voucher):
    """Excel 导出走的是同一段区间 —— 屏幕上是 700，文件里也得是 700。"""
    import io

    from openpyxl import load_workbook

    _open_years(db, book, [2027, 2028])
    make_voucher("2026-06-10", dr=[("1002", 100, "x")], cr=[("5001", 100, "x")],
                 posted=True)
    make_voucher("2027-06-10", dr=[("1002", 200, "x")], cr=[("5001", 200, "x")],
                 posted=True)
    make_voucher("2028-06-10", dr=[("1002", 400, "x")], cr=[("5001", 400, "x")],
                 posted=True)
    db.commit()

    first = _period(db, book, 2026, 1)
    last = _period(db, book, 2028, 12)
    r = client.get(f"/api/books/{book.id}/print/balance.xlsx", headers=auth, params={
        "period_id": last.id, "from_period_id": first.id})
    assert r.status_code == 200, r.text[:300]

    sheet = load_workbook(io.BytesIO(r.content)).active
    debits = []
    for row in sheet.iter_rows():
        values = [c.value for c in row]
        if any(str(v) == "1002" for v in values if v is not None):
            # 导出件的列序固定：A 科目编码 / B 科目名称 / C 方向 / D 期初 /
            # **E 本期借方** / F 本期贷方 / …
            debits.append(values[4])
    assert debits, "导出件里没有 1002 那一行"
    assert Decimal(str(debits[0])) == Decimal("700"), (
        f"导出件的区间借方应当是 700，实际 {debits[0]}")


def test_年度没取全时合并报错而不是算少(db, book, make_voucher):
    """fail-closed 断言：区间跨 N 个年度、科目表只覆盖 M 个（M<N）→ 报错。

    这道断言是「取表逻辑又漏了一年」的**最后一道网**：没有它，
    上面那种整段丢发生额的事故会**静默**发生 —— 返回的表自洽、接口 200、
    数字却少一截，而没有任何人会发现。
    """
    from app.ledger.engine import LedgerEngine
    from app.ledger.queries import subject_tree

    _open_years(db, book, [2027, 2028])
    make_voucher("2027-06-10", dr=[("1002", 200, "x")], cr=[("5001", 200, "x")],
                 posted=True)
    db.commit()

    first = _period(db, book, 2026, 1)
    last = _period(db, book, 2028, 12)

    # **故意**只取首尾两年的科目表（模拟「取表逻辑漏了中间一年」）
    incomplete = subject_tree(db, book.id, 2026) + subject_tree(db, book.id, 2028)
    assert {s.year for s in incomplete} == {2026, 2028}

    eng = LedgerEngine.for_ledger(db, book.id)
    with pytest.raises(ValueError) as exc:
        eng.merge_by_code(incomplete, last.ord, first.ord)
    message = str(exc.value)
    assert "2027" in message, f"报错要点名缺的是哪一年：{message}"
    assert "科目表" in message, f"报错要说清缺的是科目表：{message}"


def test_区间内没有科目表的空年度不误伤(db, book, make_voucher):
    """只有「那一年确实有科目表」才报错。

    区间跨到一年**压根没建过科目表**的空年度时，那一年既没有分录也没有
    期初快照，本来就没有数可丢 —— 那种情况报错只会挡住一次本来正常的查询。
    """
    from app.accounting.periods import ensure_periods
    from app.ledger.engine import LedgerEngine
    from app.ledger.queries import subject_tree

    # 2027 年只有期间、**没有科目表**
    ensure_periods(db, book, 2027)
    db.commit()

    first = _period(db, book, 2026, 1)
    last = _period(db, book, 2027, 12)

    subjects = subject_tree(db, book.id, 2026)
    eng = LedgerEngine.for_ledger(db, book.id)
    merged = eng.merge_by_code(subjects, last.ord, first.ord)
    assert merged, "空年度不该让正常的区间查询报错"
