"""总账导出的区间必须与页面一致（审计 P3-08）。

导出接口原来只认 ``period_id``，把 from/to 都当成同一期 —— 而前端页面
显示的是「起始月 ~ 结束月」这一段。于是屏幕上是一页、下载下来是全账，
数字对不上，用户只会以为导出坏了。

断言的是**导出文件里的实际内容**，不是「参数有没有传」：只查参数的话，
后端把 from_period_id 收下却不用，测试照样绿。

**总账换成用友版式之后，导出件也换了形状**：一个一级科目一张工作表，
表体逐月展开（月 / 日 / 凭证号数 / 摘要 / 借方 / 贷方 / 方向 / 余额），
「区间」不再有一行专门写它 —— 整年就是整年，区间体现在**月份列**上。
所以下面这几条断言改看月份列与文件名，区间口径一条没松。
"""
from __future__ import annotations

import io

from openpyxl import load_workbook

#: 用友的「累　　计」标签中间是**两个全角空格**（U+3000）。
CUM_LABEL = "累　　计"


def _export(client, auth, book, **params):
    q = "&".join(f"{k}={v}" for k, v in params.items())
    r = client.get(f"/api/books/{book.id}/print/general-ledger.xlsx?{q}", headers=auth)
    assert r.status_code == 200, r.text
    return load_workbook(io.BytesIO(r.content))


def _export_raw(client, auth, book, **params):
    """要连响应头一起看时用这个（文件名上带着区间）。"""
    q = "&".join(f"{k}={v}" for k, v in params.items())
    return client.get(f"/api/books/{book.id}/print/general-ledger.xlsx?{q}", headers=auth)


def _sheet(wb, code):
    """按科目编码找那张工作表 —— 一个一级科目一张表。

    工作表名是「{科目名}({编码})」（与用友的「科目: 库存现金(1001)」同源）。
    """
    for name in wb.sheetnames:
        if f"({code})" in name:
            return wb[name]
    raise AssertionError(f"导出件里没有 {code} 那张表：{wb.sheetnames}")


def _data_rows(ws):
    """表体行：从第 5 行起，到页脚之前。

    页脚那一行是「核算单位: …」，摘要列（D 列）是空的 —— 就用它当终止标记。
    """
    out = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[3] is None:
            break
        out.append(row)
    return out


def _months(ws):
    """表体里出现过的月份（按出现顺序，重复的也留着）。"""
    return [r[0] for r in _data_rows(ws) if r[0]]


def test_一个一级科目一张工作表(client, auth, db, book, periods, subjects):
    """导出件是**一册总账**：每个一级科目一张表，顺序跟科目表一致。

    以前是一张「一行一个科目」的余额表 —— 那张纸不是总账。
    """
    wb = _export(client, auth, book, period_id=periods[6].id)
    names = wb.sheetnames
    assert len(names) > 1, f"一个一级科目一张表，至少得有多张：{names}"
    assert "库存现金(1001)" in names, names
    assert "银行存款(1002)" in names, names


def test_指定区间时只导区间内的月份(db, client, auth, book, periods, subjects, make_voucher):
    """1 月与 7 月都有凭证，导出 3~5 月不该看到这两个月。"""
    make_voucher("2026-01-10", dr=[("1002", 100, "x")], cr=[("5001", 100, "x")], posted=True)
    make_voucher("2026-07-10", dr=[("1002", 500, "x")], cr=[("5001", 500, "x")], posted=True)
    db.expire_all()

    wb = _export(client, auth, book,
                 period_id=periods[5].id, from_period_id=periods[3].id)
    ws = _sheet(wb, "1002")
    # 每个月两行（本月合计 + 累计），月列两行都写同一个月份。
    assert _months(ws) == ["03", "03", "04", "04", "05", "05"], _months(ws)


def test_文件名写明的区间与传入的一致(client, auth, db, book, periods, subjects):
    """从**文件名**能看出导出的到底是什么区间。

    用友的总账表体里没有「区间」这一行 —— 整年就是整年，区间体现在月份列上。
    所以区间写不进表体就写进文件名（``总账_2026-03-2026-05.xlsx``）。
    以前是表头第二行写着「{from} 至 {to}」，现在那张纸是逐格照用友复刻的，
    那一行不能自己加 —— 区间信息一个字都不能少，只是换个地方写。
    """
    single = _export_raw(client, auth, book,
                          period_id=periods[5].id, from_period_id=periods[5].id)
    assert single.status_code == 200, single.text
    assert "2026-05.xlsx" in single.headers["content-disposition"], \
        single.headers["content-disposition"]

    multi = _export_raw(client, auth, book,
                         period_id=periods[5].id, from_period_id=periods[3].id)
    assert multi.status_code == 200, multi.text
    assert "2026-03-2026-05.xlsx" in multi.headers["content-disposition"], \
        multi.headers["content-disposition"]


def test_区间外的凭证不计入本期发生额(db, client, auth, book, periods, subjects, make_voucher):
    """5 月单月导出：1 月那笔不能出现在 5 月的「本月合计」里。

    这才是「导出的数字对不对」——总账的「本月合计」行按区间内的分录汇总，
    区间被忽略时 1 月的钱会被算进 5 月。
    """
    make_voucher("2026-01-10", dr=[("1002", 11111, "x")], cr=[("5001", 11111, "x")], posted=True)
    make_voucher("2026-05-10", dr=[("1002", 22222, "x")], cr=[("5001", 22222, "x")], posted=True)
    db.expire_all()

    wb = _export(client, auth, book,
                 period_id=periods[5].id, from_period_id=periods[5].id)
    ws = _sheet(wb, "1002")
    rows = _data_rows(ws)
    # 单月区间不是整年，标签退回「期初余额 / 本月合计 / 累计 / 期末余额」。
    assert [r[3] for r in rows] == ["期初余额", "本月合计", CUM_LABEL, "期末余额"], \
        [r[3] for r in rows]

    total = next(r for r in rows if r[3] == "本月合计")
    period_debit = float(total[4] or 0)     # 借方
    assert period_debit == 22222.0, (
        f"5 月借方应为 5 月那笔 22222，实际 {period_debit} —— 区间外的分录被算进来了")


def test_不传起始期间时退化成单期(client, auth, db, book, periods, subjects):
    """老调用方只传 period_id，行为不能变。"""
    wb = _export(client, auth, book, period_id=periods[6].id)
    ws = _sheet(wb, "1001")
    assert _months(ws) == ["06", "06"], _months(ws)


def test_不存在的起始期间返回404(client, auth, db, book, periods):
    r = client.get(f"/api/books/{book.id}/print/general-ledger.xlsx"
                   f"?period_id={periods[6].id}&from_period_id=999999", headers=auth)
    assert r.status_code == 404, f"不存在的起始期间应 404，实际 {r.status_code}"


def test_起始晚于结束时返回400而不是自动纠正(client, auth, db, book, periods, subjects):
    """**倒置区间必须报错，不能悄悄交换起止期。**

    以前这里是「起止选反了就交换一下继续跑」：同一组参数在余额表导出、
    总账打印、明细账导出上都是 400「起始期间不能晚于结束期间」，唯独总账
    Excel 导出给一份文件 —— 同一件事，三种结果。

    用户若在页面上看到过那句报错，改走导出却拿到文件，会以为页面报错是多余
    的；反过来，先导出、再在页面上核对时，两边对不上。
    """
    r = client.get(f"/api/books/{book.id}/print/general-ledger.xlsx"
                   f"?period_id={periods[2].id}&from_period_id={periods[8].id}",
                   headers=auth)
    assert r.status_code == 400, f"倒置区间应当回 400，实际 {r.status_code}：{r.text[:200]}"
    assert "起始期间不能晚于结束期间" in r.text, r.text[:200]
