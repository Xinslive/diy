"""总账（用友式）的行结构、数字算法与打印版式。

**为什么这组用例盯的是「用友逐格对齐」**：公司现在用的是用友，迁到本系统
之后，用户与审计方认的是用友那张纸。总账一旦与用友对不上，迁移就是一次
「数字打架」—— 而数字对不上的原因通常不是算错，是**口径/版式不同**
（用友的总账是「一个一级科目一份、按月展开」，不是「一行一个科目」）。

所以这里断言的不只是「有数」，还有：

* 行型与行序（上年结转 / 本月合计 / 累计 / 本年累计 / 结转下年）；
* 每行的**字段占用**（累计行不写方向与余额，期初行不写借贷方）；
* 「期初 + 累计借 − 累计贷 = 期末」这个恒等式；
* 一级科目要把它下面所有明细科目的分录**上卷**进来；
* 打印件上的页号是「当前页-该科目总页数」，两行表头，8 列。
"""
from __future__ import annotations

import io
import re
from decimal import Decimal

from openpyxl import load_workbook

from app.ledger import queries


def _ledger(db, book, subject_code, periods, from_month=1, to_month=12):
    subject = next(s for s in db.query(
        queries.Subject).filter(
            queries.Subject.book_id == book.id,
            queries.Subject.code == subject_code).all())
    return queries.general_ledger(
        db, book.id, subject, periods[from_month], periods[to_month], True)


def _rows_of(payload, kind):
    return [r for r in payload["rows"] if r["kind"] == kind]


def _load_ledger_baseline():
    """按文件路径加载 scripts/set_ledger_baseline.py。

    不走 ``from scripts import set_ledger_baseline``：``scripts/`` 没有
    ``__init__.py``，它是一堆独立运维脚本而不是一个包（同 test_db_backup_restore）。
    """
    import importlib.util
    from pathlib import Path

    root = Path(__file__).resolve().parents[1]
    spec = importlib.util.spec_from_file_location(
        "cwq_set_ledger_baseline", root / "scripts" / "set_ledger_baseline.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# --------------------------------------------------------------------- 行结构


def test_整年区间是上年结转加每月两行加结转下年(db, book, periods, subjects):
    """一整年 = 1 + 12×2 + 1 = 26 行，与用友样张一致。"""
    payload = _ledger(db, book, "1001", periods)
    kinds = [r["kind"] for r in payload["rows"]]
    assert kinds[0] == "OPENING", f"第一行应是上年结转，实际 {kinds[0]}"
    assert kinds[-1] == "CLOSING", f"最后一行应是结转下年，实际 {kinds[-1]}"
    assert len(payload["rows"]) == 26, f"整年总账应是 26 行，实际 {len(payload['rows'])}"
    assert kinds.count("MONTH_TOTAL") == 12
    assert kinds.count("YTD_TOTAL") == 12


def test_行标签与用友一致(db, book, periods, subjects):
    payload = _ledger(db, book, "1001", periods)
    summaries = [r["summary"] for r in payload["rows"]]
    assert summaries[0] == "上年结转", summaries[0]
    assert summaries[-1] == "结转下年", summaries[-1]
    # 1~11 月的累计行写「累　　计」，12 月那行写「本年累计」——用友的固定写法。
    # 「累　　计」中间是**两个全角空格**（U+3000）：用友的标签按四字宽排版，
    # 普通空格只有半字宽，与样张并排看字距对不上。
    assert summaries[2] == "累　　计", repr(summaries[2])
    assert summaries[-2] == "本年累计", summaries[-2]
    assert summaries[1] == "本月合计" and summaries[3] == "本月合计"


def test_累计行不写方向与余额_期初行不写借贷方(db, book, periods, subjects):
    """**字段占用也照用友来**：累计行只有借/贷两列有数，方向与余额是空的。

    用友的样张里，方向与余额只出现在「上月结转」「本月合计」「结转下年」
    三种行上。填了就是多出来的格子 —— 与用友并排放两张纸，一眼就能看出
    不是同一张。
    """
    payload = _ledger(db, book, "1001", periods)
    for row in payload["rows"]:
        if row["kind"] == "YTD_TOTAL":
            assert row["direction"] == "", "累计行不该有方向"
            assert row["balance"] is None, "累计行不该有余额"
            assert row["debit"] is not None and row["credit"] is not None
        if row["kind"] == "OPENING":
            assert row["debit"] is None and row["credit"] is None, "期初行不该有借贷方"
            assert row["balance"] is not None
        if row["kind"] == "MONTH_TOTAL":
            assert row["direction"] and row["balance"] is not None


def test_月列在本月合计与累计两行都写同一个月份(db, book, periods, subjects):
    payload = _ledger(db, book, "1001", periods)
    by_month = {}
    for row in payload["rows"]:
        if row["kind"] in ("MONTH_TOTAL", "YTD_TOTAL"):
            by_month.setdefault(row["month"], []).append(row["month"])
    assert by_month["01"] == ["01", "01"], by_month.get("01")
    assert by_month["12"] == ["12", "12"], by_month.get("12")
    # 期初行与结转行不写月份
    assert payload["rows"][0]["month"] == ""
    assert payload["rows"][-1]["month"] == ""


# --------------------------------------------------------------------- 数字算法


def test_本月合计等于该月发生额_累计逐月累加(
        db, book, periods, subjects, make_voucher):
    """造两笔凭证，逐月核对。"""
    make_voucher("2026-01-10", dr=[("1001", 100, "x")], cr=[("5001", 100, "x")],
                 posted=True)
    make_voucher("2026-03-15", dr=[("1001", 50, "y")], cr=[("5001", 50, "y")],
                 posted=True)
    db.expire_all()

    payload = _ledger(db, book, "1001", periods)
    months = {r["month"]: r for r in _rows_of(payload, "MONTH_TOTAL")}
    assert months["01"]["debit"] == Decimal("100")
    assert months["01"]["credit"] == Decimal("0")
    assert months["02"]["debit"] == Decimal("0"), "2 月没有凭证"
    assert months["03"]["debit"] == Decimal("50")

    cums = {r["month"]: r for r in _rows_of(payload, "YTD_TOTAL")}
    assert cums["01"]["debit"] == Decimal("100")
    assert cums["02"]["debit"] == Decimal("100"), "2 月累计 = 1 月 + 2 月"
    assert cums["03"]["debit"] == Decimal("150")
    assert cums["12"]["debit"] == Decimal("150"), "本年累计 = 全年"


def test_期初加累计等于期末(db, book, periods, subjects, make_voucher):
    """**「期初 + 累计借 − 累计贷 = 期末」在每一行都成立。**

    这是总账这张纸的自洽性：期初行、月度合计行、结转行说的是同一个账，
    三处对不上就是印错了一页。区间不是整年时累计是区间累计，恒等式照样成立。
    """
    make_voucher("2026-01-10", dr=[("1001", 100, "x")], cr=[("5001", 100, "x")],
                 posted=True)
    make_voucher("2026-03-15", dr=[("1001", 50, "y")], cr=[("5001", 50, "y")],
                 posted=True)
    db.expire_all()

    payload = _ledger(db, book, "1001", periods)
    opening = _rows_of(payload, "OPENING")[0]
    closing = _rows_of(payload, "CLOSING")[0]
    assert opening["balance"] == Decimal("0"), "没有期初余额时是 0"
    assert closing["balance"] == Decimal("150")

    cum = _rows_of(payload, "YTD_TOTAL")[-1]
    assert (opening["balance"] + cum["debit"] - cum["credit"]) == closing["balance"]


def test_月末余额逐月递推(db, book, periods, subjects, make_voucher):
    make_voucher("2026-01-10", dr=[("1001", 100, "x")], cr=[("5001", 100, "x")],
                 posted=True)
    make_voucher("2026-03-15", dr=[("1001", 50, "y")], cr=[("5001", 50, "y")],
                 posted=True)
    db.expire_all()

    payload = _ledger(db, book, "1001", periods)
    months = _rows_of(payload, "MONTH_TOTAL")
    assert months[0]["balance"] == Decimal("100")
    assert months[1]["balance"] == Decimal("100"), "2 月没动，余额不变"
    assert months[2]["balance"] == Decimal("150")


def test_一级科目把明细科目的分录上卷进来(
        db, book, periods, subjects, make_voucher):
    """分录挂在 110101（短期投资的明细）上，一级科目 1101 的总账要算到它。

    **这是总账与明细账的根本区别**：用友的总账是「总分类账」，一个一级科目
    一份，把它下面所有明细科目的流水汇总进去。少了这一层，一级科目的总账
    就是空的 —— 而它在用友那张纸上是有数的。
    """
    make_voucher("2026-02-10", dr=[("110101", 800, "x")], cr=[("5001", 800, "x")],
                 posted=True)
    db.expire_all()

    payload = _ledger(db, book, "1101", periods)
    feb = [r for r in _rows_of(payload, "MONTH_TOTAL") if r["month"] == "02"][0]
    assert feb["debit"] == Decimal("800"), "明细科目的分录没上卷到一级科目"
    closing = _rows_of(payload, "CLOSING")[0]
    assert closing["balance"] == Decimal("800")


def test_贷方余额的方向写贷(db, book, periods, subjects, make_voucher):
    """方向看**符号**，不看科目属性（引擎口径，见 display_side）。"""
    make_voucher("2026-01-10", dr=[("5001", 100, "x")], cr=[("1001", 100, "x")],
                 posted=True)
    db.expire_all()

    payload = _ledger(db, book, "1001", periods)
    closing = _rows_of(payload, "CLOSING")[0]
    assert closing["direction"] == "贷", closing["direction"]
    assert closing["balance"] == Decimal("100"), "贷方余额存绝对值"


def test_非整年区间退回期初余额与期末余额(db, book, periods, subjects):
    """区间不是整年时，标签**不许**写「上年结转」/「结转下年」。

    写「上年结转」而实际是三月一日的期初，是拿标签骗人 —— 审计一眼就看出来。
    """
    payload = _ledger(db, book, "1001", periods, from_month=3, to_month=6)
    summaries = [r["summary"] for r in payload["rows"]]
    assert summaries[0] == "期初余额", summaries[0]
    assert summaries[-1] == "期末余额", summaries[-1]
    assert "本年累计" not in summaries
    assert len(payload["rows"]) == 1 + 4 * 2 + 1


# --------------------------------------------------------------------- 打印


def test_打印件是每页一份的用友版式(client, auth, book, periods):
    html = client.get(f"/api/books/{book.id}/print/general-ledger.html",
                       headers=auth,
                       params={"period_id": periods[12].id,
                               "from_period_id": periods[1].id}).text
    assert "总账" in html
    # 两行表头：年份横跨「月 + 日」
    assert re.search(r'colspan="2"[^>]*class="lb-year"', html), "缺少年份表头"
    assert ">月<" in html and ">日<" in html
    # 八个列头。**标签不带全角空格**（用户自己在设计器里调出来的那份
    # 基线模板就是这个写法：摘要 / 借方 / 贷方 / 余额）。
    for label in ("凭证号数", "摘要", "借方", "贷方", "余额"):
        assert label in html, f"缺列头：{label}"
    # 「方向」列窄到 3%，表头必须竖排（方在上、向在下）——用友的固定写法
    assert "<th rowspan=\"2\">方<br>向</th>" in html, "方向列头没有竖排"
    # 页眉：页号 / 本币名称
    assert "本币名称: 人民币" in html
    assert re.search(r"页号:\s+1-1", html), "页号应是「当前页-该科目总页数」"
    # **右上角那个账簿标记不印**（用户 2026-10-22 批注 1 要求去掉用友的
    # 「[正式账簿]」）。字段还在模板里 —— 用户自己填回去就会照印。
    assert "[正式账簿]" not in html, "右上角不该再印用友的账簿标记"
    assert '<div class="lb-book">' not in html, "空的账簿标记连元素都不该输出"


def test_页号是该科目自己的总页数(client, auth, book, periods, subjects):
    """**每个一级科目各自从 1 开始编页**：1-1 = 只有一页，这是第一页。

    用友的页号不是「第几册-第几页」，而是「当前页-该科目总页数」。
    """
    html = client.get(f"/api/books/{book.id}/print/general-ledger.html",
                       headers=auth,
                       params={"period_id": periods[12].id,
                               "from_period_id": periods[1].id}).text
    # 一个一级科目一份总账，每份第一页都是 1-1
    assert html.count("页号:  1-1") >= 1
    # 不该出现「册号-页号」那种 2-1
    assert not re.search(r"页号:\s+2-1", html)


def test_每页的数据行数不超过模板设定(client, auth, book, periods):
    html = client.get(f"/api/books/{book.id}/print/general-ledger.html",
                       headers=auth,
                       params={"period_id": periods[12].id,
                               "from_period_id": periods[1].id}).text
    bodies = re.findall(r"<tbody>(.*?)</tbody>", html, re.S)
    assert bodies, "没有表体"
    for body in bodies:
        rows = re.findall(r"<tr>", body)
        assert len(rows) <= 26, f"一页 {len(rows)} 行，超过模板设定的 26 行"


def test_导出与打印是同一张纸(client, auth, book, periods, make_voucher):
    """Excel 导出与 HTML 打印印的是同一张纸：都是一级科目的月度台账。

    账上要有数才有东西可导 —— 整册科目都没动过时，「不打印无发生额科目」
    （默认开）会把它们全跳掉，导出件就只剩一张空白页（那也是对的）。
    """
    make_voucher("2026-01-10", dr=[("1001", 100, "x")], cr=[("5001", 100, "x")],
                 posted=True)
    r = client.get(f"/api/books/{book.id}/print/general-ledger.xlsx",
                   headers=auth,
                   params={"period_id": periods[12].id,
                           "from_period_id": periods[1].id})
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    # 一级科目一张表
    assert len(wb.sheetnames) >= 1
    ws = wb[wb.sheetnames[0]]
    texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    blob = " ".join(texts)
    assert "总账" in blob
    assert "上年结转" in blob, "导出件里没有「上年结转」行"
    assert "本月合计" in blob
    assert "本年累计" in blob
    assert "结转下年" in blob
    # 表头两行：年份 + 月/日
    assert any(t.endswith("年") and len(t) == 5 for t in texts), "表头缺年份"


# --------------------------------------------------------------------- 跳过无发生额


def _ledger_html(client, auth, book, periods, **params):
    query = {"period_id": periods[12].id, "from_period_id": periods[1].id}
    query.update(params)
    r = client.get(f"/api/books/{book.id}/print/general-ledger.html",
                   headers=auth, params=query)
    assert r.status_code == 200, r.text
    return r.text


def test_默认跳过无发生额的一级科目(client, auth, book, periods, make_voucher):
    """「不打印无发生额科目」**默认开**（用户 2026-10-22 批注 2）。

    一册总账是几十个一级科目，每一个都是 26 行；期间内没动过的那些印出来
    就是几十页全是零的表 —— 那不是账簿，是废纸。
    """
    make_voucher("2026-01-10", dr=[("1001", 100, "x")], cr=[("5001", 100, "x")],
                 posted=True)
    html = _ledger_html(client, auth, book, periods)
    assert "库存现金" in html, "有发生额的科目必须在"
    assert "库存商品" not in html, "没动过的一级科目不该印出来"


def test_关掉开关就全都印(client, auth, book, periods, make_voucher):
    """开关是**打印时**的参数（打印预览弹窗里那个），不是模板里的字段。"""
    make_voucher("2026-01-10", dr=[("1001", 100, "x")], cr=[("5001", 100, "x")],
                 posted=True)
    html = _ledger_html(client, auth, book, periods, skip_empty="false")
    assert "库存现金" in html
    assert "库存商品" in html, "关了开关之后，没动过的科目也该印"


# --------------------------------------------------------------------- 版式基线


def test_默认版式是横向七毫米行高():
    """基线模板是**用户自己在设计器里调出来的那份**（2026-10-22）。

    横向 A4 + 行高 7mm + 摘要列占 42% —— 横向才放得下用友那种长摘要。
    代码里的默认值与库里的基线行是同一份内容，两边不该各说各话。
    """
    from app.printing.template import DEFAULT_TEMPLATES, validate_template

    tpl = validate_template("LEDGER", DEFAULT_TEMPLATES["LEDGER"])
    assert tpl["orientation"] == "landscape", "基线版式是横向"
    assert tpl["ledger_book"]["row_height_mm"] == 7.0
    assert tpl["margins"] == {"top": 12.0, "right": 20.0,
                              "bottom": 12.0, "left": 20.0}
    # 列标签不带全角空格（用户那份的写法）
    labels = [c["label"] for c in tpl["columns"]]
    assert labels == ["月", "日", "凭证号数", "摘要", "借方", "贷方", "方向", "余额"]
    # 页脚中段是空的 —— 用户把「制表: {preparer}」清掉了
    assert tpl["ledger_book"]["footer_center"] == ""


def test_基线脚本把账套模板提为基线(db, book):
    """**把用户调出来的那份模板提为内置默认基线**（用户 2026-10-22）。

    基线是模板三层归属里最外面那一层（``book_id IS NULL AND user_id IS NULL``），
    与凭证范围那两份同一层。提为基线之后它对**所有账套**都生效。

    基线行的内容取自代码里的 ``DEFAULT_TEMPLATES`` —— 代码默认值与库里的
    基线行是同一份内容，两边不会各说各话（改版式只改一处，另一处由脚本同步）。
    """
    from app.models.printing import PrintTemplate
    from app.printing.template import DEFAULT_TEMPLATES, validate_template
    from sqlalchemy import select

    script = _load_ledger_baseline()
    for scope, name in script.BASELINE_NAMES.items():
        # 先造一份「用户在设计器里调出来的账套模板」
        row = PrintTemplate(book_id=book.id, user_id=None, scope=scope,
                            name="新默认模板", paper="A4", content={},
                            is_default=False, remark="")
        db.add(row)
        db.flush()

        script.apply_scope(db, scope, apply=True)
        db.flush()

        base = db.scalar(select(PrintTemplate).where(
            PrintTemplate.scope == scope,
            PrintTemplate.book_id.is_(None),
            PrintTemplate.user_id.is_(None)))
        assert base is not None, f"{scope} 没有基线模板行"
        assert base.is_default, f"{scope} 的基线模板不是默认模板"
        assert base.name == name
        assert base.id == row.id, "应该是把用户那一行提为基线，而不是另起一行"
        saved = validate_template(scope, base.content or {})
        expect = validate_template(scope, DEFAULT_TEMPLATES[scope])
        for key in ("orientation", "margins", "title", "columns", "ledger_book"):
            assert saved[key] == expect[key], f"{scope} 的 {key} 与代码默认不一致"


# --------------------------------------------------------------------- 副标题


def test_标题下面是期间副标题(client, auth, book, periods):
    """「标题下面显示期间的副标题」在总账上是**活的**（用户 2026-10-22）。

    账簿渲染器自己拼页眉，以前整段没有副标题 —— 于是模板设计器左栏那个
    「副标题」输入框对总账/明细账是死的：填什么都不出东西。
    """
    html = _ledger_html(client, auth, book, periods)
    assert re.search(r'<div class="lb-subtitle">\s*2026-01 至 2026-12\s*</div>', html), \
        "标题下面没有期间副标题"


def test_同一期只写一个期间(client, auth, book, periods):
    """起止是同一期时只写一个 —— 「X 至 X」是把同一件事说了两遍。"""
    html = client.get(f"/api/books/{book.id}/print/general-ledger.html",
                      headers=auth,
                      params={"period_id": periods[12].id}).text
    assert '<div class="lb-subtitle">2026-12</div>' in html, \
        "同一期不该写成「起 至 止」"


def test_标题下面没有横线(client, auth, book, periods):
    """用户 2026-10-22：「把主标题下面的横线去掉吧」。

    用友的样张上标题下有线，我们抄过来之后看着像一条多余的框线。
    """
    html = _ledger_html(client, auth, book, periods)
    assert "border-bottom" not in html, "标题下面还有那条横线"


# --------------------------------------------------------------------- 分页


def test_一行放不下就分页_页码与纸上页数一致(client, auth, book, periods,
                                                make_voucher):
    """**分页要按纸面算**，不能按模板里写的行数算。

    基线模板是横向 A4（可印高度 186mm）+ 行高 7mm：26 行放不下一页。硬塞的
    结果是浏览器把它接到下一张物理页上 —— 第一页写着「1-1」，第二张物理页
    还是「1-1」。按纸面算出来的行数分页，页码才与纸上对得上。
    """
    from app.printing.template import validate_template, ledger_book_rows_per_page

    tpl = validate_template("LEDGER", {})
    per_page = ledger_book_rows_per_page(tpl)
    assert per_page < 26, "横向 + 7mm 行高放不下 26 行，分页必须比它少"

    # 要有发生额：否则整册被「跳过无发生额科目」清空，拿到的是那张空白页
    make_voucher("2026-01-10", dr=[("1001", 100, "x")], cr=[("5001", 100, "x")],
                 posted=True)
    html = _ledger_html(client, auth, book, periods)
    # 一个整年的总账是 26 行 → 按实放行数分两页
    assert "页号:  1-2" in html, "第一页的页号不是「1-2」"
    assert "页号:  2-2" in html, "第二页的页号不是「2-2」"
    assert "页号:  1-1" not in html, "不该还按模板写的 26 行当成一页"


def test_一个逻辑页就是一张物理纸(client, auth, book, periods):
    """每个册页单位从 1 重新编页 —— 前提是**一个逻辑页 = 一张物理纸**。

    没有强制分页时，两个册页单位会接在同一张纸上连续印（第一张装得下时，
    第二张直接跟在后面），纸上第二页写着「2-2」、物理上却是第一张纸的
    下半截。
    """
    html = _ledger_html(client, auth, book, periods)
    assert "break-after:page" in html, "没有强制分页，逻辑页与物理页会对不上"
    assert "page-break-after:always" in html


# --------------------------------------------------------------------- 册页顺序


def test_册页单位按科目编码排(db, book):
    """一册账的目录按**科目编码数字**排，不是科目表的分类顺序。

    科目表按会计分类排（资产 → 所有者权益 → 负债 → 损益），照那个顺序打出来
    「本年利润(4103)」会排在「应付账款(2202)」前面 —— 翻账的人是**按编码
    找**的，这个目录看着是乱的（用户 2026-10-22）。
    """
    from app.ledger import queries

    subjects = queries.general_ledger_subjects(db, book.id, book.start_year)
    codes = [s.code for s in subjects]
    assert codes, "账套里一个一级科目都没有"
    assert codes == sorted(codes, key=lambda c: int(c)), \
        f"册页单位没按科目编码排：{codes[:8]}…"
    if "2202" in codes and "4103" in codes:
        assert codes.index("2202") < codes.index("4103"), \
            "应付账款(2202)应该排在本年利润(4103)前面"


def test_明细账的册页单位同样按科目编码排(db, book):
    """明细账与总账是**同一册结构**，目录顺序也得一样。"""
    from app.ledger import queries

    groups = queries.detail_ledger_groups(db, book.id, book.start_year)
    codes = [g["subject"].code for g in groups]
    assert codes, "账套里一个一级科目都没有"
    assert codes == sorted(codes, key=lambda c: int(c)), \
        f"册页单位没按科目编码排：{codes[:8]}…"
