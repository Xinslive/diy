"""审计 §3.26 的回归：报表数值在**四条通道**上是同一个数。

审计 §3.26 的原症状
--------------------
「报表单元格：屏幕按两位小数舍入，导出文件保留原始 Decimal → 屏幕与文件
不是同一个数」。报表里用除法公式（1/3 = 0.333…），**屏幕上写 0.33**，
**导出的 Excel 里是 0.3333333333333333333333333333**，两处对不上，且没有
任何提示。报表是给审计/外部看的正式件，这种对不上不能留。

修法，以及本组用例盯住的不变量
------------------------------
口径（**两位小数**）定在**唯一的生产者**上：``app/reports/evaluate.py``。
页面、打印预览、Excel 导出、MCP 全部读它同一个输出 —— 所以四处必然一致。

本组用例用**同一份 payload** 走完四条通道，**逐字**比对。只要有一条通道
自己再舍入一次（或压根不舍入），比对就会红。

**为什么不是「改屏幕还是改导出」**
导出侧的两位是**既有事实**：Excel 单元格格式 ``#,##0.00``、HTML 打印
``_fmt`` 一律两位、MCP 一直是两位。把屏幕改成 28 位小数反而离另外三条通道
更远 —— 而且那不是给人读的数字，还会把「MCP 与界面逐字一致」这条已有
不变量拆散（``tests/test_mcp_caliber.py`` 盯着）。

**为什么不在每条通道各写一份舍入**
审计 §4.5：同一个量只允许一个权威实现。所以舍入规则只有一份，在
``app/core/money.py``；报表、MCP 共用它。
"""
from __future__ import annotations

import io
import uuid
from decimal import Decimal

import openpyxl
import pytest

from app.core.money import money_str, money_str_or_text
from app.models import ReportCell, ReportCol, ReportRow, ReportTemplate
from app.printing import excel as xl
from app.printing import preview as pv
from app.printing.template import DEFAULT_TEMPLATES, render_table_page, validate_template
from app.reports.evaluate import evaluate


# --------------------------------------------------------------- 造一张最小报表


def _one_cell_template(db, book, formula: str) -> int:
    """造一张**只有一格数字**的报表模板（列号 1、行号 1）。

    只有一格是有意的：Excel 里于是只有一个数字，「文件里的那个数」没有
    歧义，比对才成立。
    """
    tag = uuid.uuid4().hex[:8]
    t = ReportTemplate(code=f"CAL_{tag}", book_id=book.id, name=f"口径用例{tag}",
                       type="CUSTOM", remark="")
    db.add(t)
    db.flush()
    db.add(ReportRow(template_id=t.id, row_no=1, label="比值",
                     indent=0, bold=False, hidden=False, style={}))
    db.add(ReportCol(template_id=t.id, col_no=1, label="金额",
                     kind="AMOUNT", hidden=False, width=120))
    db.add(ReportCell(template_id=t.id, row_no=1, col_no=1, formula=formula))
    db.commit()
    db.refresh(t)
    return t.id


def _cell_value(payload: dict) -> str:
    """取 payload 里那一格的数值（字符串形态）。"""
    cells = payload["rows"][0]["cells"]
    return cells[0]["value"]


def _excel_cells(payload: dict) -> tuple[list, list]:
    """把报表导成 xlsx，再把里面的**数值格**与**文本格**分开读回来。"""
    buf = xl.report_excel(payload, book_name="口径用例")
    book = openpyxl.load_workbook(io.BytesIO(buf), data_only=True)
    nums: list = []
    texts: list = []
    for ws in book.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, bool) or c.value in (None, ""):
                    continue
                if isinstance(c.value, (int, float)):
                    nums.append(c.value)
                else:
                    texts.append(c.value)
    return nums, texts


# ------------------------------------------------------------------- 单元层


def test_除法结果按两位小数给出():
    """1/3 的十进制展开是 28 位；对外只给两位。

    这是整条链的**起点**：口径定在这里，四条通道才可能一致。
    """
    third = Decimal(1) / Decimal(3)
    assert len(str(third)) > 10, "前置条件：1/3 本来就该有很长的小数位"
    assert money_str_or_text(third) == "0.33"


def test_不是金额的文本原样带出去():
    """公式留下的文本**不能**被改成 0.00 —— 那会把「这里有错」变成「这里是零」。"""
    assert money_str_or_text("科目不存在") == "科目不存在"
    assert money_str_or_text(None) == ""
    assert money_str_or_text("") == ""


def test_舍入规则只有一份():
    """MCP 通道与报表通道必须是**同一个**实现（审计 §4.5）。

    ``app/mcp/caliber.py`` 的 ``money_str`` 现在只是 ``app/core/money.py``
    的转手 —— 两边对同一个数必须给出同一个字符串。
    """
    from app.mcp.caliber import money_str as mcp_money_str

    cases = [
        Decimal("0.33"),
        Decimal(1) / Decimal(3),          # 1/3：28 位小数，除法不收窄
        Decimal("12800"),
        Decimal("-0.005"),                # 恰好落在半分上：舍入规则在这里才显形
        Decimal("999999999999999999.99"),  # 大额：float 装不下
    ]
    for d in cases:
        assert mcp_money_str(d) == money_str(d), d


# ------------------------------------------------------------------ 端到端


def test_四条通道给出同一个数(db, book, periods):
    """**同一份 payload**，页面 / Excel / HTML 打印 / MCP 逐字一致。

    这是 §3.26 的正面回答：屏幕与文件不再是两个数。
    """
    from app.mcp.tools.analysis import _cell_money

    tid = _one_cell_template(db, book, "1/3")
    payload = evaluate(db, tid, book, periods[1], False)

    # ① 页面 / MCP 读的那一份
    screen = _cell_value(payload)
    assert screen == "0.33", screen

    # ② Excel 导出
    nums, texts = _excel_cells(payload)
    assert len(nums) == 1, f"这张表只有一格数字，导出里却有 {len(nums)} 个数值格：{nums}"
    assert Decimal(str(nums[0])) == Decimal("0.33"), nums
    # 文件里不能藏着 28 位那个数（那就是「屏幕与文件不是同一个数」）
    assert all("0.3333333333333333" not in str(n) for n in nums + texts), nums + texts

    # ③ HTML 打印（走产品自己的那一段：打印列由报表列生成）
    from app.api.printing import _report_print_col

    tpl = dict(DEFAULT_TEMPLATES["REPORT"])
    tpl["columns"] = [_report_print_col(c) for c in payload["cols"]]
    tpl = validate_template("REPORT", tpl)
    html = render_table_page(tpl, {}, pv.report_rows(payload), 1, 1)
    assert "0.33" in html
    assert "0.3333333333333333" not in html

    # ④ MCP 出参
    assert _cell_money(payload["rows"][0]["cells"][0])["value"] == "0.33"


def test_口径说明随数据一起走(db, book, periods):
    """审计 §3.26 要求「导出 quantize 到两位**并在界面注明**」。

    说明随 payload 走（前端直接显示它），但**不能**混进 ``notices`` ——
    notices 是「本期有什么特殊情况」的通道（没结转时必须是空数组）。
    """
    tid = _one_cell_template(db, book, "1/3")
    payload = evaluate(db, tid, book, periods[1], False)

    scale = payload["display_scale"]
    assert scale["decimals"] == 2
    assert scale["note"], "口径说明不能是空串 —— 界面要能把它显示出来"
    # notices 的契约没被动过
    assert isinstance(payload["notices"], list)


def test_大额在Excel里不丢分(db, book, periods):
    """超过 Excel 15 位有效数字的金额，文件里**不能是另一个数**。

    openpyxl 写 Decimal 走 ``"%.16g"``（先转 float），实测
    ``Decimal("99999999999999.99")`` 会被写成 ``99999999999999.98`` ——
    差一分，而屏幕上是准的。那就还是「屏幕与文件不是同一个数」。

    超出 Excel 数值格能力的，改写成**文本**：数字一个字节不差，
    代价是那格在 Excel 里左对齐（带「数字存为文本」的绿三角）——
    **看得见**的错位，好过一个看不出差别的错数。
    """
    big = "99999999999999.99"
    tid = _one_cell_template(db, book, big)
    payload = evaluate(db, tid, book, periods[1], False)

    screen = _cell_value(payload)
    assert screen == big, screen

    found, texts = _excel_cells(payload)
    # 文本格里是**精确的数字**（带千分位，与屏幕上的显示逐字相同）
    assert "99,999,999,999,999.99" in texts, texts
    # 而且不能同时存在一个「差一分」的数值格 —— 那正是原来的症状
    assert not found, f"文件里还有一个数值格 {found}：它和文本格不是同一个数"


def test_两位以内的大额仍然是数值格(db, book, periods):
    """Excel 放得下的金额必须**写成数值**（不是文本）。

    否则一列金额全是文本，Excel 自己求和会把它们全跳过去 —— 那比丢一分
    更糟：合计看起来对，其实少算了一整列。
    """
    ok = "1234567890123.45"          # 15 位有效数字，Excel 放得下
    tid = _one_cell_template(db, book, ok)
    payload = evaluate(db, tid, book, periods[1], False)
    assert _cell_value(payload) == ok

    buf = xl.report_excel(payload, book_name="口径用例")
    book_x = openpyxl.load_workbook(io.BytesIO(buf), data_only=True)
    cells = [c for ws in book_x.worksheets for row in ws.iter_rows() for c in row
             if c.value not in (None, "")]
    numeric = [c for c in cells if isinstance(c.value, (int, float))]
    assert numeric, "放得下的金额被写成了文本 —— Excel 求和会把它跳过去"
    assert Decimal(str(numeric[0].value)) == Decimal(ok)
